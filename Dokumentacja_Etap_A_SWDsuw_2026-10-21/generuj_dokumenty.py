from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parent
DATE = '21.10.2026'
PROJECT = 'System wspomagania dyplomowania studentów uczelni wyższej (SWDsuw)'
NAVY = '1F4E78'
BLUE = 'D9EAF7'
PALE = 'F3F6F9'
ORANGE = 'FCE4D6'
GREEN = 'E2F0D9'
YELLOW = 'FFF2CC'
PURPLE = 'E4DFEC'
RED = 'F4CCCC'


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run('Strona ')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)


def base_doc(title: str, doc_id: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)

    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(35, 44, 54)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 10.5)]:
        style = doc.styles[style_name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)

    header = section.header.paragraphs[0]
    header.text = f'SWDSUW  |  FAZA DEFINICJI  |  {doc_id}'
    header.style = doc.styles['Caption']
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.add_run(f'{doc_id}  ·  Wersja 1.0  ·  {DATE}                                      ')
    add_page_number(footer)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(title)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 78, 121)
    p = doc.add_paragraph(subtitle)
    p.paragraph_format.space_after = Pt(18)
    p.runs[0].font.size = Pt(12)
    p.runs[0].font.color.rgb = RGBColor(91, 103, 115)

    info = doc.add_table(rows=0, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.LEFT
    info.style = 'Table Grid'
    for label, value in [
        ('Projekt', PROJECT),
        ('Identyfikator', doc_id),
        ('Wersja / data', f'1.0 / {DATE}'),
        ('Autor', 'Zespół projektowy — nazwiska do uzupełnienia'),
        ('Status', 'Projekt roboczy do przeglądu Zamawiającego'),
    ]:
        cells = info.add_row().cells
        cells[0].text = label
        cells[1].text = value
        shade(cells[0], BLUE)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    hist = doc.add_table(rows=1, cols=4)
    hist.style = 'Table Grid'
    for cell, text in zip(hist.rows[0].cells, ['Data', 'Wersja', 'Opis', 'Autor']):
        cell.text = text
        shade(cell, NAVY)
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
    row = hist.add_row().cells
    for cell, text in zip(row, [DATE, '1.0', 'Pierwsza uporządkowana wersja artefaktu', 'Zespół projektowy']):
        cell.text = text
    doc.add_page_break()
    return doc


def h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str, style=None) -> None:
    p = doc.add_paragraph(style=style)
    p.add_run(text)


def bullet(doc: Document, text: str) -> None:
    para(doc, text, 'List Bullet')


def numbered(doc: Document, text: str) -> None:
    para(doc, text, 'List Number')


def table(doc: Document, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(t.rows[0])
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = str(value)
        shade(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.5)
    for i, row_values in enumerate(rows):
        row = t.add_row()
        for cell, value in zip(row.cells, row_values):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2:
                shade(cell, PALE)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph()
    return t


def callout(doc, title, body, color=YELLOW):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    shade(t.cell(0, 0), color)
    p = t.cell(0, 0).paragraphs[0]
    r = p.add_run(title + '\n')
    r.bold = True
    p.add_run(body)
    doc.add_paragraph()


CLIENT_PARAGRAPHS = [
    'Każdy student kończący etap edukacji (dyplomant) musi wykonać i obronić pracę końcową, zwaną pracą dyplomową. Uczelnia składa się z kilku wydziałów. Na każdym z nich studenci mogą kształcić się na różnych specjalizacjach. Informacja o wydziale i specjalizacji jest istotna przy sporządzaniu sprawozdań dotyczących obronionych prac. Praca dyplomowa jest realizowana na koniec każdego etapu studiów: pierwszego, drugiego lub trzeciego stopnia. Uczelnia planuje w przyszłości rozszerzyć ofertę o nowe rodzaje studiów, które również będą kończyć się pracą dyplomową, np. studia podyplomowe. Każdy student ma jednoznaczny, 10-znakowy numer albumu.',
    'Student chcący ukończyć dany rodzaj studiów wybiera temat pracy dyplomowej i związanego z nim opiekuna, zwanego promotorem. Jeden temat jest prowadzony przez jednego promotora — pracownika uczelni ze stopniem doktora, doktora habilitowanego lub profesora.',
    'Uczelnia zachęca do pisania prac indywidualnych, ale dopuszcza także pracę dwóch lub trzech studentów. Jednej pracy nie może przygotowywać więcej niż trzech studentów.',
    'Przed dopuszczeniem do obrony praca jest sprawdzana w elektronicznym systemie „Plagiat” pod kątem podobieństwa do dokumentów z jego bazy i zasobów Internetu. Z opisu Zamawiającego wynika, że pozytywna weryfikacja pozwala przekazać pracę do obrony.',
    'Praca podlega recenzji jednego lub kilku pracowników uczelni. Każdy recenzent ocenia pracę w skali od 2 do 6. Pracę ocenia również promotor. Ocena promotora i recenzentów musi być przypisana osobno każdemu autorowi, a nie wyłącznie całej grupie.',
    'Do pracy przypisuje się słowa kluczowe z wcześniej zdefiniowanego zbioru. Umożliwiają one określenie tematyki i wyszukiwanie prac. Każda praca ma mieć co najmniej jedno słowo kluczowe.',
    'Obrona odbywa się w wyznaczonym dniu przed komisją składającą się z trzech lub maksymalnie czterech członków: promotora, jednego lub dwóch recenzentów i przewodniczącego. Studenci przedstawiają prezentację pracy. Każdy student otrzymuje pytania, które są protokołowane. Komisja wystawia każdemu studentowi osobną ocenę końcową.',
    'Komisja ogłasza wynik ukończenia studiów, obliczany według algorytmu określonego w statucie uczelni, oraz sporządza wymagane raporty.',
    'Uczelnia chce usprawnić obsługę prac dyplomowych i powiązanych procesów przez opracowanie systemu informatycznego wspierającego tę obsługę.',
]

FIRST_STAGE = [
    'Zarządzanie informacjami o pracach dyplomowych, przydziale recenzentów, komisjach i egzaminach dyplomowych.',
    'Wspomaganie negocjacji dotyczących wyboru tematu pracy dyplomowej.',
    'Bezpieczne przechowywanie prac i przekazywanie ich do systemu antyplagiatowego; udostępnianie prac osobom uprawnionym oraz zarządzanie informacjami o pracach obronionych.',
    'Szybkie wyszukiwanie prac według tematyki, wydziału, instytutu, specjalizacji lub promotora.',
    'Wspomaganie obrony w zakresie pytań, oceny końcowej, ocen cząstkowych członków komisji i składu komisji.',
    'Zarządzanie raportami systemu, w tym raportami z obron.',
]


def build_client_problem():
    doc = base_doc('Przedstawienie problemu — dokument Zamawiającego', 'A1', 'Uporządkowany zapis źródłowego przedstawienia problemu')
    h(doc, '1. Cel działalności i problem', 1)
    para(doc, 'Uczelnia prowadzi proces dyplomowania studentów wielu wydziałów, specjalizacji i poziomów studiów. Wraz ze wzrostem liczby rodzajów studiów, autorów prac i uczestników procesu potrzebny jest system wspierający obsługę prac od wyboru tematu do obrony i raportowania.')
    h(doc, '2. Opis sytuacji Zamawiającego', 1)
    for text in CLIENT_PARAGRAPHS[:3]:
        para(doc, text)
    h(doc, '3. Sprawdzenie pracy, recenzje i obrona', 1)
    for text in CLIENT_PARAGRAPHS[3:8]:
        para(doc, text)
    h(doc, '4. Oczekiwany cel systemu', 1)
    para(doc, CLIENT_PARAGRAPHS[8])
    h(doc, '5. Zakres pierwszego etapu', 1)
    para(doc, 'Pierwszy etap prac ma doprowadzić do projektu systemu zapewniającego:')
    for text in FIRST_STAGE:
        bullet(doc, text)
    h(doc, '6. Uwagi redakcyjne i punkty do potwierdzenia', 1)
    callout(doc, 'Charakter dokumentu', 'Dokument porządkuje treść źródłową. Nie uzupełnia brakujących reguł biznesowych jako faktów. Poniższe kwestie wymagają potwierdzenia przez Zamawiającego.')
    for text in [
        'Kryterium pozytywnej kontroli podobieństwa, procedura rozbieżności i treść zwrotu z systemu „Plagiat”.',
        'Liczba pytań: opis mówi o 3–4 pytaniach od każdego członka komisji, a następnie o wspólnej średniej 3–4 pytań. Należy ustalić, czy liczba dotyczy całej obrony, czy każdego członka komisji.',
        'Reguły wyznaczania oceny końcowej i wyniku ukończenia studiów ze statutu uczelni; nie podano algorytmu.',
        'Format, odbiorcy, częstotliwość i uprawnienia do raportów.',
        'Czy numer albumu składa się z liter i cyfr; czy ma być unikalny w obrębie całej uczelni.',
    ]:
        bullet(doc, text)
    return doc


STAKEHOLDERS = [
    ('Zamawiający / władze uczelni', 'Właściciel procesu i budżetu; zatwierdza zakres, reguły i kryteria odbioru.'),
    ('Student / dyplomant', 'Wybiera temat, przygotowuje pracę, przekazuje ją do oceny i broni.'),
    ('Promotor', 'Prowadzi jeden temat, ocenia każdego autora indywidualnie i uczestniczy w komisji.'),
    ('Recenzent', 'Przygotowuje recenzję i indywidualną ocenę autora.'),
    ('Przewodniczący komisji', 'Koordynuje obronę, zapewnia protokołowanie pytań i kończenie procesu.'),
    ('Komisja dyplomowa', 'Prowadzi obronę i wystawia indywidualną ocenę końcową.'),
    ('Administrator / pracownik dziekanatu', 'Kandydat na rolę operacyjną do potwierdzenia; utrzymuje dane organizacyjne i raporty.'),
    ('Zewnętrzny system antyplagiatowy', 'System „Plagiat” wskazany w opisie; zakres techniczny i kontrakt integracyjny nie są opisane.'),
]

PROCESSES = [
    ('P1', 'Wybór i uzgodnienie tematu', 'Student wybiera temat; system wspiera negocjacje i wiąże temat z jednym promotorem.'),
    ('P2', 'Rejestracja autorów i pracy', 'Praca jest związana z jednym, dwoma lub trzema autorami oraz ich numerami albumów.'),
    ('P3', 'Przekazanie pracy do weryfikacji', 'Praca jest bezpiecznie przechowywana i przekazywana do systemu „Plagiat”; dalsze kryteria wyniku wymagają decyzji.'),
    ('P4', 'Recenzowanie i ocena', 'Promotor i jeden lub kilku recenzentów oceniają każdego autora oddzielnie w skali 2–6.'),
    ('P5', 'Przygotowanie i przebieg obrony', 'Tworzy się komisję 3–4-osobową; studenci prezentują pracę, pytania są protokołowane, a wyniki indywidualne.'),
    ('P6', 'Wyznaczenie wyniku i raportowanie', 'Komisja ogłasza wynik według algorytmu statutowego i tworzy wymagane raporty.'),
    ('P7', 'Wyszukiwanie prac', 'Uprawnieni użytkownicy wyszukują prace według tematu, wydziału, instytutu, specjalizacji lub promotora.'),
]

OPEN_QUESTIONS = [
    ('OQ-01', 'Kryterium antyplagiatowe', 'Jaki próg podobieństwa/rezultat pozwala na dopuszczenie do obrony? Jak obsłużyć wynik niejednoznaczny lub błąd integracji?'),
    ('OQ-02', 'Pytania komisji', 'Czy 3–4 pytania są łącznie na studenta, czy od każdego członka komisji? Co oznacza „wspólna średnia ocena końcowa”?'),
    ('OQ-03', 'Oceny', 'Jakie są dozwolone wartości skali 2–6 (np. połówki)? Jaki dokładnie algorytm statutowy wylicza ocenę końcową i wynik ukończenia studiów?'),
    ('OQ-04', 'Recenzenci', 'Czy liczba recenzentów jest ograniczona? Kto i według jakich reguł ich przydziela?'),
    ('OQ-05', 'Użytkownicy i dostęp', 'Jakie role mogą widzieć pracę, recenzje, oceny, protokół i raporty? Jak rozstrzygać konflikt ról promotora/recenzenta/przewodniczącego?'),
    ('OQ-06', 'Numer albumu', 'Jakie 10 znaków są dopuszczalne i kto nadaje numer? Czy numer jest globalnie unikalny?'),
    ('OQ-07', 'Słowa kluczowe', 'Kto zatwierdza słownik słów kluczowych, ich synonimy i zmiany?'),
    ('OQ-08', 'System „Plagiat”', 'Jaka jest nazwa dostawcy, API/protokół, format pliku, limity, zasady retencji i wymagane zgody?'),
    ('OQ-09', 'Raporty', 'Jakie raporty i formaty są wymagane, przez kogo zatwierdzane i jak długo przechowywane?'),
    ('OQ-10', 'Zakres organizacyjny', 'Jaka uczelnia i które wydziały będą pierwszym wdrożeniem? Jak identyfikować instytuty i stopnie studiów?'),
]


def build_contractor_problem():
    doc = base_doc('Przedstawienie problemu — dokument Wykonawcy', 'A2', 'Analiza potrzeb i kontekstu rozwiązania SWDsuw')
    h(doc, '1. Cel dokumentu i źródła', 1)
    para(doc, 'Celem dokumentu jest zinterpretowanie przedstawienia Zamawiającego jako wspólnego opisu problemu biznesowego, kontekstu, aktorów i ograniczeń. Podstawowym źródłem jest „PZ SWD przedstawienie problemu.docx”. Szablon WAT dla przedstawienia problemu posłużył do uporządkowania sekcji: cel działalności, opis organizacji, aktorzy, procesy biznesowe i ograniczenia.')
    callout(doc, 'Granica źródeł', 'Dokument MSIoPS jest wyłącznie przykładem struktury. Jego aktorzy, funkcje i wymagania nie są wymaganiami SWDsuw.')
    h(doc, '2. Opis problemu', 1)
    para(doc, 'Uczelnia musi koordynować wieloetapowy proces dyplomowania: identyfikację studenta i pracy, wybór tematu i promotora, weryfikację antyplagiatową, recenzje, obronę, ocenę każdego autora i raportowanie. Dane zależą od wydziału, specjalizacji i poziomu studiów. Opis zakłada rozwój oferty studiów, więc model procesu powinien uwzględniać przyszłe typy studiów bez utraty spójności danych.')
    para(doc, 'Zasadniczym problemem jest brak wskazanego, zintegrowanego sposobu zarządzania tymi informacjami i procesami w jednym miejscu. Zamawiający oczekuje systemu wspomagającego obsługę, nie zastępującego merytorycznych decyzji promotora, recenzentów, komisji ani algorytmu statutowego.')
    h(doc, '3. Cel biznesowy', 1)
    for text in [
        'Zapewnić spójne zarządzanie informacją o pracach, tematach, autorach, promotorach, recenzjach, komisjach, obronach i raportach.',
        'Umożliwić indywidualne oceny i wynik dla każdego autora pracy wieloosobowej.',
        'Ułatwić uprawnione wyszukiwanie prac według atrybutów domenowych wskazanych przez Zamawiającego.',
        'Ograniczyć ręczne przekazywanie dokumentów i zapewnić kontrolowany obieg do weryfikacji antyplagiatowej.',
        'Zachować możliwość obsługi nowych rodzajów studiów kończących się pracą dyplomową.',
    ]:
        bullet(doc, text)
    h(doc, '4. Aktorzy biznesowi i interesariusze', 1)
    table(doc, ['Aktor / interesariusz', 'Oczekiwanie / rola'], STAKEHOLDERS)
    h(doc, '5. Procesy biznesowe w zakresie pierwszego etapu', 1)
    table(doc, ['ID', 'Proces', 'Interpretacja na podstawie źródła'], PROCESSES)
    h(doc, '6. Granica systemu i zakres', 1)
    para(doc, 'System SWDsuw obejmuje organizację danych i przebieg obsługi opisanych procesów. System „Plagiat” jest systemem zewnętrznym, którego usługi wykorzystuje proces kontroli; nie ma podstaw do zakładania, że dostawca lub interfejs techniczny są już wybrani.')
    h(doc, 'W zakresie wynikającym ze źródła', 2)
    for text in FIRST_STAGE:
        bullet(doc, text)
    h(doc, 'Poza ustalonym zakresem / do decyzji', 2)
    for text in [
        'Zakup, wymiana lub wdrożenie systemu antyplagiatowego — opis wymaga integracji z istniejącym systemem, a nie jego zastąpienia.',
        'Szczegóły finansowania, wdrożenia na konkretnych wydziałach, migracji danych i kanałów powiadomień — nie zostały określone.',
        'Technologie, architektura, baza danych, protokół integracji i platforma wdrożenia — pozostają poza decyzjami w fazie definicji.',
        'Wymagania prawne, retencja prac, zgody i procedura odwoławcza — wymagają uzupełnienia przez uczelnię.',
    ]:
        bullet(doc, text)
    h(doc, '7. Fakty i założenia projektowe', 1)
    table(doc, ['Rodzaj', 'Treść'], [
        ('Fakt źródłowy', 'Praca ma 1–3 autorów; każdy ma 10-znakowy numer albumu.'),
        ('Fakt źródłowy', 'Jeden temat jest przypisany do jednego promotora o stopniu co najmniej doktora.'),
        ('Fakt źródłowy', 'Każda praca ma co najmniej jedno słowo kluczowe.'),
        ('Fakt źródłowy', 'Recenzje i oceny muszą być przypisane oddzielnie każdemu autorowi.'),
        ('Fakt źródłowy', 'Komisja liczy 3–4 członków; dokładny przebieg punktacji wymaga wyjaśnienia.'),
        ('Założenie do zatwierdzenia', 'System rozdziela uprawnienia co najmniej na studenta/dyplomanta, promotora, recenzenta, członka/przewodniczącego komisji i rolę administracyjną.'),
        ('Założenie do zatwierdzenia', 'Zewnętrzny system antyplagiatowy zwraca identyfikowalny rezultat kontroli, który można powiązać z pracą i wersją pliku.'),
    ])
    h(doc, '8. Ograniczenia', 1)
    for text in [
        'Zakres dziedzinowy wynika z zasad uczelni i statutu; algorytm oceny nie został dostarczony.',
        'Skala ocen recenzji i promotora: 2–6; zasady stosowania ocen pośrednich nie zostały określone.',
        'Maksymalnie trzech studentów może być autorami jednej pracy.',
        'Prace zawierają dane osobowe i dokumenty akademickie wymagające dostępu wyłącznie dla osób uprawnionych.',
        'Nazwy firmy, uczelni, wydziałów, autorów i docelowej platformy nie zostały podane w źródle SWDsuw.',
    ]:
        bullet(doc, text)
    h(doc, '9. Ryzyka i zależności', 1)
    table(doc, ['Ryzyko / zależność', 'Wpływ', 'Działanie w definicji'], [
        ('Nieznany algorytm statutowy', 'Błędne wyniki końcowe i raporty.', 'Pozyskać zatwierdzony algorytm oraz przykłady testowe od uczelni.'),
        ('Nieznane API/zasady systemu „Plagiat”', 'Brak możliwości bezpiecznego lub terminowego przekazania prac.', 'Warsztat z administratorem/dostawcą; uzgodnić format, statusy, retencję i błędy.'),
        ('Niespójny opis pytań na obronie', 'Różne interpretacje obliczania oceny.', 'Rozstrzygnąć OQ-02 przed specyfikacją i odbiorem.'),
        ('Brak matrycy uprawnień', 'Ujawnienie pracy, recenzji albo ocen.', 'Zdefiniować role, zakres i audyt dostępu z inspektorem ochrony danych.'),
        ('Zmienne typy studiów', 'Założenia ograniczą rozbudowę oferty.', 'Potwierdzić, które dane są konfigurowalne, a które stałe.'),
    ])
    h(doc, '10. Rejestr kwestii wymagających decyzji', 1)
    table(doc, ['ID', 'Kwestia', 'Pytanie do Zamawiającego'], OPEN_QUESTIONS)
    return doc


def build_vision():
    doc = base_doc('Wizja projektu SWDsuw', 'A3', 'Dokument wizji zgodny ze strukturą szablonu fazy definicji')
    h(doc, '1. Wstęp', 1)
    h(doc, '1.1 Cel', 2)
    para(doc, 'Dokument uzgadnia problem, cel, grupy użytkowników, wysokopoziomowy zakres i ograniczenia systemu SWDsuw. Nie jest specyfikacją techniczną ani decyzją o konkretnych technologiach.')
    h(doc, '1.2 Zakres', 2)
    para(doc, 'Wizja obejmuje system wspomagający procesy dyplomowania od zarządzania tematem i pracą przez kontrolę antyplagiatową i recenzje po obronę, ocenę i raportowanie. Zakres wymaga zatwierdzenia przez Zamawiającego.')
    h(doc, '1.3 Definicje, akronimy i skróty', 2)
    para(doc, 'SWDsuw — System wspomagania dyplomowania studentów uczelni wyższej. Pozostałe pojęcia są opisane w artefakcie A6 „Słownik pojęć SWDsuw”.')
    h(doc, '1.4 Dokumenty powiązane', 2)
    table(doc, ['Dokument', 'Rola', 'Źródło / uwaga'], [
        ('PZ SWD przedstawienie problemu.docx', 'Źródło potrzeb Zamawiającego.', 'Katalog „Dokumenty wejściowe”.'),
        ('Przedstawienie problemu — dokument Wykonawcy', 'Analiza problemu, granicy i kwestii do decyzji.', 'Artefakt A2.'),
        ('Słownik projektu SWDsuw', 'Jednoznaczna terminologia domenowa.', 'Artefakt A6.'),
        ('RUP_Project.mpp', 'Szablon struktury zadań i ról RUP.', 'Daty w pliku są przykładowe (2005); nie są harmonogramem SWDsuw.'),
    ])
    h(doc, '1.5 Organizacja dokumentu', 2)
    para(doc, 'Kolejne sekcje opisują pozycjonowanie rozwiązania, interesariuszy i użytkowników, ich potrzeby, ograniczenia, wymagania jakościowe i priorytety. Braki oznaczono jako kwestie do potwierdzenia, a nie jako uzgodnione wymagania.')

    h(doc, '2. Pozycjonowanie projektu', 1)
    h(doc, '2.1 Cele biznesowe', 2)
    for text in [
        'Spójnie koordynować dane pracy dyplomowej, jej autora lub autorów, temat, promotora, recenzentów, komisję, obronę i raporty.',
        'Umożliwić odrębną ocenę każdego autora pracy wieloosobowej.',
        'Chronić poufność prac oraz udostępniać je wyłącznie osobom uprawnionym.',
        'Usprawnić wyszukiwanie prac i raportowanie według danych wskazanych przez uczelnię.',
        'Obsłużyć nowe rodzaje studiów kończących się pracą bez przebudowy całego procesu.',
    ]:
        bullet(doc, text)
    h(doc, '2.2 Opis problemu / potrzeby', 2)
    table(doc, ['Pole szablonu', 'Treść'], [
        ('Problem', 'Informacje i działania związane z dyplomowaniem wymagają skoordynowania od wyboru tematu do wyniku obrony i raportów.'),
        ('Dotyczy', 'Dyplomantów, promotorów, recenzentów, komisji, administracji uczelni i osób uprawnionych do raportowania.'),
        ('Znaczenie', 'Nieuporządkowany obieg może utrudniać terminową obsługę, kontrolę dostępu, przypisanie ocen do autorów i wyszukiwanie prac.'),
        ('Pomyślne rozwiązanie', 'Jeden spójny proces obsługi pracy, jawne relacje między autorami a ocenami, kontrolowany obieg dokumentu i raportowalne wyniki.'),
    ])
    h(doc, '2.3 Deklaracja pozycjonowania produktu', 2)
    callout(doc, 'Dla', 'uczelni wyższej, która prowadzi prace dyplomowe na wielu wydziałach, specjalizacjach i stopniach studiów,')
    callout(doc, 'SWDsuw', 'jest systemem wspomagającym obsługę pracy dyplomowej od wyboru tematu do obrony i raportowania. W odróżnieniu od nieuzgodnionych narzędzi cząstkowych łączy dane i przebieg procesu w jednym modelu domenowym. Obecne alternatywy i stan automatyzacji wymagają rozpoznania; nie zakłada się, że proces jest dziś realizowany w określonej technologii.')

    h(doc, '3. Interesariusze i użytkownicy', 1)
    table(doc, ['Typ', 'Opis', 'Odpowiedzialność / interes'], STAKEHOLDERS)
    h(doc, '3.1 Środowisko użytkownika', 2)
    para(doc, 'Oczekiwane środowisko obejmuje pracę studentów i pracowników uczelni w okresie przygotowania, recenzowania oraz obron. Liczba użytkowników, urządzenia, lokalizacja, dostępność, integracja tożsamościowa i obciążenie nie zostały określone; muszą zostać zebrane przed wymaganiami niefunkcjonalnymi.')
    h(doc, '3.2 Profile reprezentatywne', 2)
    table(doc, ['Profil', 'Cel', 'Sukces z punktu widzenia roli', 'Dane do potwierdzenia'], [
        ('Dyplomant', 'Wybrać temat, złożyć pracę, poznać oceny i przygotować się do obrony.', 'Wszystkie wymagane kroki i statusy są czytelne; wynik dotyczy jego osoby.', 'Wymagane powiadomienia, widoczność recenzji, odwołania.'),
        ('Promotor', 'Prowadzić temat, nadzorować pracę, ocenić autorów.', 'Widoczna jest praca każdego autora i indywidualna ocena.', 'Limity liczby tematów, zamiana promotora, konflikty interesów.'),
        ('Recenzent', 'Zapoznać się z uprawnioną pracą i przygotować recenzję.', 'Może ocenić pracę i poszczególnych autorów zgodnie z regułami.', 'Szablon recenzji, terminy, możliwość wielu recenzji.'),
        ('Przewodniczący komisji', 'Zorganizować obronę i zamknąć jej protokół.', 'Wynik, pytania i skład komisji są kompletne i zatwierdzone.', 'Harmonogram, podpisy, korekta protokołu.'),
        ('Administrator / dziekanat', 'Utrzymywać dane organizacyjne i raporty.', 'Dane są spójne, a zakres uprawnień kontrolowany.', 'Rzeczywisty profil, integracje i uprawnienia.'),
    ])
    h(doc, '3.3 Potrzeby klienta i priorytety', 2)
    table(doc, ['Priorytet roboczy', 'Potrzeba', 'Status'], [
        ('MUST — ze źródła', 'Zarządzanie pracami, tematami, przydziałami, recenzjami, komisjami i egzaminami.', 'Potwierdzić szczegółowe statusy i reguły.'),
        ('MUST — ze źródła', 'Wspomaganie negocjacji wyboru tematu.', 'Zdefiniować, kto proponuje i zatwierdza temat.'),
        ('MUST — ze źródła', 'Bezpieczne przechowywanie i przekazanie pracy do systemu antyplagiatowego.', 'Uzgodnić dostawcę, protokół, retencję, wynik i obsługę błędu.'),
        ('MUST — ze źródła', 'Wyszukiwanie po temacie, wydziale, instytucie, specjalizacji i promotorze.', 'Zdefiniować role i widoczność wyników.'),
        ('MUST — ze źródła', 'Wspomaganie obrony i indywidualnych ocen.', 'Rozstrzygnąć liczbę pytań i algorytm wyniku.'),
        ('MUST — ze źródła', 'Zarządzanie raportami systemu.', 'Uzgodnić typy raportów i odbiorców.'),
    ])
    h(doc, '3.4 Alternatywy i konkurencja', 2)
    para(doc, 'Źródło nie opisuje obecnych narzędzi ani konkurencyjnych produktów. W tej wersji nie przypisuje się uczelni obiegu papierowego, arkuszy, konkretnego systemu dziekanatowego ani konkretnego procesu ręcznego. Wymagane jest rozpoznanie stanu istniejącego.')
    h(doc, '4. Ograniczenia', 1)
    for text in [
        'Praca dyplomowa ma od jednego do trzech autorów.',
        'Jeden temat ma jednego promotora; promotorem może być pracownik o stopniu doktora, doktora habilitowanego lub profesora.',
        'Numer albumu identyfikuje studenta i ma 10 znaków; do ustalenia pozostaje zestaw dozwolonych znaków.',
        'Każda praca ma co najmniej jedno słowo kluczowe ze zdefiniowanego zbioru.',
        'Recenzent i promotor oceniają każdego autora oddzielnie; opis wskazuje skalę 2–6.',
        'Komisja składa się z 3–4 osób i obejmuje promotora, jednego lub dwóch recenzentów oraz przewodniczącego.',
        'Wynik ukończenia studiów zależy od algorytmu statutu, którego nie dostarczono.',
        'Nie podejmuje się decyzji technologicznych w niniejszym dokumencie.',
    ]:
        bullet(doc, text)
    h(doc, '5. Wymagania jakościowe — do kwantyfikacji', 1)
    table(doc, ['Cecha', 'Oczekiwany kierunek', 'Co trzeba uzgodnić'], [
        ('Poufność i uprawnienia', 'Prace, recenzje, oceny i protokoły dostępne wyłącznie rolom uprawnionym.', 'Matryca ról, log audytowy, przypadki konfliktu ról.'),
        ('Integralność ocen', 'Oceny przypisane właściwemu autorowi i recenzentowi; reguły obliczeń odtwarzalne.', 'Algorytm, skala, zaokrąglenia, korekty.'),
        ('Niezawodność integracji', 'Awaria systemu antyplagiatowego nie powoduje utraty pracy ani błędnego dopuszczenia.', 'SLA, retry, statusy i ścieżka ręczna.'),
        ('Użyteczność', 'Proces jest zrozumiały dla studentów i pracowników o różnym doświadczeniu technicznym.', 'Badania użytkowników, urządzenia i dostępność.'),
        ('Wydajność / skala', 'Wyszukiwanie i obieg pozostają akceptowalne przy obciążeniu uczelni.', 'Liczby studentów, prac, jednoczesnych sesji i oczekiwany czas odpowiedzi.'),
    ])
    h(doc, '6. Priorytety', 1)
    para(doc, 'Wizja zachowuje priorytety MUST wynikające z dokumentu Zamawiającego. Nie przypisuje poziomów SHOULD/COULD do funkcji nieopisanych w źródle. Szacowanie i priorytety iteracji zatwierdza Zamawiający po warsztacie wymagań.')
    h(doc, '7. Pozostałe wymagania i standardy', 1)
    para(doc, 'Źródło nie określa platformy, przeglądarek, wymagań sprzętowych, standardów interfejsów, integracji kont uczelnianych, okresów retencji ani norm prawnych. Należy uzupełnić je po konsultacji z administratorem IT, właścicielem procesu i inspektorem ochrony danych.')
    return doc


SOFTAS_TEXT = [
    'Dyrektor firmy „Soft-AS” zlecił firmie AS-Projektowanie realizację systemu: „System wsparcia sprzedaży” (SWS). Zasadniczym celem tego systemu będzie wspomaganie pracy handlowców.',
    'Firma „Soft-AS”, zajmująca się dystrybucją materiałów biurowych, uznała, że potrzebny jest jej system informatyczny SWS wspierający pracę przedstawicieli handlowych (popularnie zwanych handlowcami), zarówno zatrudnionych bezpośrednio w firmie Soft-AS, jak również w hurtowniach z tą firmą współpracujących (hurtownie mają całkowicie niezależne magazyny). Wprowadzenie tego systemu powinno zmienić również sposób pracy magazynierów, ponieważ zamówienia będą spływały do istniejącego systemu magazynowego równomiernie przez cały dzień (optymalizując pracę osób zatrudnionych w magazynie); obecnie wszystkie zamówienia spływają z samego rana (wysyłane faksem przez handlowca za cały poprzedni dzień).',
]

NOUNS = [
    ('Dyrektor firmy „Soft-AS”', 'dyrektor', 'Interesariusz / osoba inicjująca zlecenie'),
    ('firmie AS-Projektowanie', 'firma', 'Organizacja wykonawcza; poza granicą systemu'),
    ('system „System wsparcia sprzedaży” (SWS)', 'system', 'Produkt będący przedmiotem zlecenia'),
    ('praca handlowców', 'praca', 'Cel biznesowy — wspomaganie sprzedaży'),
    ('Firma „Soft-AS”', 'firma', 'Organizacja zamawiająca i dystrybutor'),
    ('materiałów biurowych', 'materiał', 'Towar domenowy; potrzebę katalogu trzeba potwierdzić'),
    ('system informatyczny SWS', 'system', 'System docelowy'),
    ('przedstawicieli handlowych', 'przedstawiciel', 'Rola biznesowa / aktor'),
    ('handlowcami', 'handlowiec', 'Synonim przedstawiciela handlowego'),
    ('hurtowniach', 'hurtownia', 'Partner dystrybucyjny; ma własny magazyn'),
    ('firmą Soft-AS', 'firma', 'Organizacja współpracująca z hurtownią'),
    ('niezależne magazyny', 'magazyn', 'Zasób lokalny każdej hurtowni'),
    ('sposób pracy magazynierów', 'magazynier', 'Rola użytkownika i jego proces'),
    ('zamówienia', 'zamówienie', 'Główny kandydat na encję z własnym cyklem życia'),
    ('systemu magazynowego', 'system magazynowy', 'System zewnętrzny / punkt integracji'),
    ('osób zatrudnionych w magazynie', 'osoba / magazynier', 'Użytkownicy procesu magazynowego'),
    ('faksem', 'faks', 'Kanał obecnego przekazania zamówień, nie obiekt domenowy'),
    ('poprzedni dzień', 'dzień', 'Atrybut czasu opisujący bieżący obieg zamówień'),
]

VERBS = [
    ('zlecił realizację systemu', 'zlecić / realizować', 'Dyrektor Soft-AS / AS-Projektowanie', 'Zlecenie wytworzenia produktu'),
    ('będzie wspomaganie pracy handlowców', 'wspomagać', 'SWS', 'Cel systemu; działanie produktu'),
    ('zajmująca się dystrybucją', 'zajmować się / dystrybuować', 'Soft-AS', 'Kontekst biznesowy organizacji'),
    ('uznała, że potrzebny jest system', 'uznać / być potrzebnym', 'Soft-AS', 'Rozpoznanie potrzeby'),
    ('wspierający pracę przedstawicieli', 'wspierać', 'SWS', 'Zakres wsparcia użytkowników'),
    ('zatrudnionych bezpośrednio', 'zatrudniać', 'Soft-AS', 'Relacja pracownik–organizacja'),
    ('współpracujących z tą firmą', 'współpracować', 'Hurtownia / Soft-AS', 'Relacja partnerstwa'),
    ('mają całkowicie niezależne magazyny', 'mieć', 'Hurtownie', 'Własność / separacja magazynów'),
    ('powinno zmienić sposób pracy', 'zmienić', 'Wprowadzenie SWS', 'Oczekiwany skutek wdrożenia'),
    ('zamówienia będą spływały równomiernie', 'spływać', 'Zamówienia / system magazynowy', 'Zdarzenie przepływu danych'),
    ('optymalizując pracę osób', 'optymalizować', 'SWS / proces', 'Oczekiwany efekt operacyjny'),
    ('zamówienia spływają z samego rana', 'spływać', 'Obecny proces', 'Opis stanu obecnego'),
    ('wysyłane faksem przez handlowca', 'wysyłać', 'Handlowiec', 'Obecny kanał przekazania'),
]

NOUN_MARKS = [x[0] for x in NOUNS]
VERB_MARKS = ['zlecił', 'będzie wspomaganie', 'zajmująca się', 'uznała', 'jest potrzebny', 'wspierający', 'zatrudnionych', 'współpracujących', 'mają', 'powinno zmienić', 'będą spływały', 'optymalizując', 'spływają', 'wysyłane']


def add_highlighted_text(doc, text):
    p = doc.add_paragraph()
    matches = []
    for phrase in NOUN_MARKS:
        for match in re.finditer(re.escape(phrase), text, re.IGNORECASE):
            matches.append((match.start(), match.end(), 'noun', phrase))
    for phrase in VERB_MARKS:
        for match in re.finditer(re.escape(phrase), text, re.IGNORECASE):
            matches.append((match.start(), match.end(), 'verb', phrase))
    matches.sort(key=lambda item: (item[0], -(item[1] - item[0]), 0 if item[2] == 'noun' else 1))
    chosen = []
    end = -1
    for item in matches:
        if item[0] >= end:
            chosen.append(item)
            end = item[1]
    cursor = 0
    for start, stop, kind, _ in chosen:
        if start > cursor:
            p.add_run(text[cursor:start])
        run = p.add_run(text[start:stop])
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW if kind == 'noun' else WD_COLOR_INDEX.BRIGHT_GREEN
        cursor = stop
    if cursor < len(text):
        p.add_run(text[cursor:])
    p.paragraph_format.line_spacing = 1.18
    return p


def build_linguistic_analysis():
    doc = base_doc('Analiza lingwistyczna — tekst Soft-AS / SWS', 'A5', 'Kolorowanie źródła, grupy czasownikowe i rzeczownikowe, kandydaci na obiekty')
    callout(doc, 'Rozdzielenie zakresów', 'Tekst Soft-AS/SWS dostarczony w poleceniu różni się od dokumentacji projektu SWDsuw. Niniejsza analiza dotyczy wyłącznie zadanego fragmentu Soft-AS/SWS. Kandydatów z tego ćwiczenia nie przenosi się do słownika SWDsuw.')
    h(doc, '1. Klucz kolorów i metoda', 1)
    table(doc, ['Kolor podświetlenia', 'Znaczenie'], [
        ('Żółty', 'Grupa rzeczownikowa / uczestnik, organizacja, obiekt lub pojęcie.'),
        ('Zielony', 'Grupa czasownikowa / działanie, relacja lub oczekiwany skutek.'),
    ])
    para(doc, 'Kolorowanie wskazuje fragmenty analizowane, a nie gotową klasę programistyczną. Najpierw ustalono grupy składniowe, następnie oceniono ich znaczenie i cykl życia w domenie. Same rzeczowniki nie są automatycznie obiektami systemu.')
    h(doc, '2. Pokolorowany tekst źródłowy', 1)
    add_highlighted_text(doc, SOFTAS_TEXT[0])
    add_highlighted_text(doc, SOFTAS_TEXT[1])
    h(doc, '3. Grupy czasownikowe', 1)
    table(doc, ['Grupa w tekście', 'Bezokolicznik / rdzeń', 'Wykonawca / kontekst', 'Znaczenie biznesowe'], VERBS)
    h(doc, '4. Grupy rzeczownikowe', 1)
    table(doc, ['Grupa w tekście', 'Głowa grupy', 'Klasyfikacja semantyczna'], NOUNS)
    h(doc, '5. Wynik analizy: kandydaci na obiekty projektowe', 1)
    para(doc, 'Poniższe kandydatury wynikają z połączenia grup rzeczownikowych z działaniami, interesami i odrębnymi cyklami życia. Nie są bezwarunkową listą klas programu. Właściciel domeny musi zdecydować, które elementy system ma rzeczywiście ewidencjonować.')
    h(doc, '5.1 Zamówienie — kandydat na główną encję', 2)
    para(doc, 'Uzasadnienie: zamówienie jest tworzone przez handlowca, trafia do systemu magazynowego, ma moment i sposób przekazania, a równomierny wpływ jest mierzalnym celem. Kandydat powinien mieć identyfikator, źródło (Soft-AS lub hurtownia), pozycje, czas utworzenia/przekazania i status. Pozycje zamówienia oraz statusy nie są literalnie opisane — należy je potwierdzić, a nie dopisywać jako fakt.')
    h(doc, '5.2 Handlowiec i magazynier — aktorzy / role', 2)
    para(doc, 'Handlowiec (przedstawiciel handlowy) składa lub przekazuje zamówienie; magazynier realizuje jego obsługę. Z tekstu wynika, że handlowcy mogą być zatrudnieni w Soft-AS lub w hurtowniach, zatem rola użytkownika nie powinna być utożsamiana z jedną organizacją.')
    h(doc, '5.3 Hurtownia, Soft-AS i magazyn — organizacje i zasoby', 2)
    para(doc, 'Soft-AS i hurtownia są odrębnymi podmiotami biznesowymi. Każda hurtownia posiada niezależny magazyn; model musi zachować tę separację, jeżeli zakres potwierdzi wspólną obsługę podmiotów. „Magazyn” jest kandydatem na obiekt domenowy tylko wtedy, gdy system SWS ma odwzorowywać powiązanie zamówienia z konkretną lokalizacją lub stanem magazynowym.')
    h(doc, '5.4 Towar / materiał biurowy — kandydat warunkowy', 2)
    para(doc, 'Materiały biurowe określają branżę Soft-AS, ale tekst nie mówi, że SWS będzie prowadzić katalog, ceny, zapasy ani pozycje zamówienia. Towar jest kandydatem do dalszego wywiadu, nie wymaganiem potwierdzonym.')
    h(doc, '5.5 System magazynowy, faks i strumień zamówień', 2)
    para(doc, 'Istniejący system magazynowy jest systemem zewnętrznym i granicą integracji. Faks jest kanałem obecnego procesu, nie musi być obiektem w SWS. „Równomierny przepływ zamówień” jest celem/procesem i miernikiem operacyjnym, a nie automatycznie encją domenową.')
    h(doc, '6. Relacje do weryfikacji', 1)
    for item in [
        'Handlowiec składa zamówienie dla klienta; w przekazanym tekście klient nie jest nazwany — należy potwierdzić, czy jest potrzebny jako osobny obiekt.',
        'Zamówienie jest kierowane do magazynu podmiotu, który je złożył; potwierdzić regułę routingu.',
        'SWS przekazuje zamówienie do istniejącego systemu magazynowego; potwierdzić, czy integracja jest automatyczna, czy kolejkowana.',
        'Cel „równomiernie przez cały dzień” wymaga określenia miary (np. rozkład czasu wpływu) i ograniczeń systemu magazynowego.',
    ]:
        bullet(doc, item)
    h(doc, '7. Zakres i ograniczenia analizy', 1)
    para(doc, 'Analiza opiera się wyłącznie na dwóch akapitach dostarczonych przez Zamawiającego. Nie ma specyfikacji zamówienia, klienta, cen, katalogu, zapasów, zwrotów, ról administracyjnych ani API. Brak tych pojęć nie oznacza, że system ma je obsłużyć.')
    return doc


GLOSSARY = [
    ('SWDsuw', 'System wspomagania dyplomowania studentów uczelni wyższej — nazwa systemu opisanego w materiale wejściowym.', 'Nazwa projektu'),
    ('Uczelnia', 'Organizacja prowadząca studia, dyplomowanie i obrony; nazwa konkretnej uczelni nie została podana.', 'Organizacja'),
    ('Wydział', 'Jednostka organizacyjna uczelni, do której przypisany jest student lub praca.', 'Organizacja'),
    ('Instytut', 'Jednostka organizacyjna używana w kryteriach wyszukiwania; relacja do wydziału wymaga potwierdzenia.', 'Organizacja'),
    ('Specjalizacja', 'Kierunek/specjalność kształcenia studenta, wykorzystywana do klasyfikacji i raportowania prac.', 'Klasyfikacja studiów'),
    ('Stopień studiów', 'Poziom kształcenia wskazany w opisie: pierwszy, drugi lub trzeci stopień; w przyszłości mogą dojść inne rodzaje studiów.', 'Klasyfikacja studiów'),
    ('Student', 'Osoba kształcąca się na uczelni, która może wybrać temat, przygotować pracę i przystąpić do obrony.', 'Rola'),
    ('Dyplomant', 'Student kończący etap edukacji i przygotowujący się do obrony pracy dyplomowej.', 'Rola'),
    ('Numer albumu', 'Jednoznaczny, 10-znakowy identyfikator studenta w uczelni. Zestaw dopuszczalnych znaków nie został określony.', 'Identyfikator domenowy'),
    ('Praca dyplomowa', 'Praca końcowa studenta lub grupy studentów, podlegająca weryfikacji, recenzji i obronie.', 'Encja'),
    ('Autor pracy', 'Student przypisany do pracy jako jej autor; każdy autor może otrzymać odrębną ocenę.', 'Rola / relacja'),
    ('Grupa autorów', 'Od jednego do trzech studentów przygotowujących tę samą pracę.', 'Relacja'),
    ('Temat pracy', 'Temat wybrany do realizacji przez dyplomanta i powiązany z jednym promotorem.', 'Encja / pojęcie'),
    ('Promotor', 'Pracownik uczelni ze stopniem doktora, doktora habilitowanego lub profesora, prowadzący temat i oceniający autorów.', 'Rola'),
    ('Recenzent', 'Pracownik uczelni, który przygotowuje recenzję i indywidualną ocenę autora pracy.', 'Rola'),
    ('Recenzja', 'Ocena merytoryczna pracy wykonywana przez jednego lub kilku recenzentów.', 'Encja'),
    ('Ocena recenzenta', 'Ocena w skali od 2 do 6, przypisana osobno autorowi; dopuszczalność ocen połówkowych wymaga decyzji.', 'Wartość'),
    ('Ocena promotora', 'Ocena pracy wystawiana przez promotora, przypisana osobno każdemu autorowi.', 'Wartość'),
    ('Słowo kluczowe', 'Pojęcie ze zdefiniowanego przez uczelnię zbioru, przypisane pracy do określenia tematyki i wyszukiwania.', 'Klasyfikacja'),
    ('Weryfikacja antyplagiatowa', 'Sprawdzenie pracy w elektronicznym systemie „Plagiat” względem dokumentów w jego bazie i zasobów Internetu.', 'Proces'),
    ('System „Plagiat”', 'Zewnętrzny system kontroli podobieństwa wskazany w opisie Zamawiającego; dostawca, interfejs i kryteria wyniku nie są określone.', 'System zewnętrzny'),
    ('Dopuszczenie do obrony', 'Decyzja/stadium procesu pozwalające studentowi przystąpić do obrony po wymaganej weryfikacji; kryterium pozytywnego wyniku wymaga uzgodnienia.', 'Status procesu'),
    ('Obrona', 'Egzamin, podczas którego autor lub autorzy przedstawiają pracę i odpowiadają na protokołowane pytania komisji.', 'Proces'),
    ('Komisja dyplomowa', 'Zespół liczący trzy lub maksymalnie cztery osoby: promotora, jednego lub dwóch recenzentów i przewodniczącego.', 'Rola zbiorowa'),
    ('Przewodniczący komisji', 'Członek komisji kierujący jej pracami; szczegółowe obowiązki proceduralne wymagają doprecyzowania.', 'Rola'),
    ('Pytanie obrony', 'Pytanie zadane studentowi podczas obrony i protokołowane; liczba i sposób agregacji ocen są niespójnie opisane i wymagają decyzji.', 'Zdarzenie'),
    ('Ocena końcowa', 'Ocena wystawiana osobno każdemu studentowi po obronie; algorytm wyznaczenia określa statut uczelni.', 'Wartość'),
    ('Wynik ukończenia studiów', 'Wynik obliczany zgodnie z algorytmem wskazanym w statucie uczelni.', 'Wartość'),
    ('Raport z obrony', 'Wymagany dokument/raport sporządzany po obronie; format, zawartość i odbiorcy wymagają decyzji.', 'Artefakt biznesowy'),
]


def build_glossary():
    doc = base_doc('Słownik pojęć projektu SWDsuw', 'A6', 'Wspólne słownictwo domenowe Zamawiającego i Wykonawcy')
    callout(doc, 'Zasada słownika', 'Słownik obejmuje wyłącznie pojęcia domeny dyplomowania. Nie zawiera technologii, baz danych, interfejsów API ani nazw bibliotek. Tam, gdzie źródło jest niejednoznaczne, definicja jest oznaczona jako wymagająca uzgodnienia.')
    h(doc, '1. Pojęcia domenowe', 1)
    table(doc, ['Nazwa pojęcia', 'Znaczenie', 'Rodzaj'], GLOSSARY)
    h(doc, '2. Zasady używania terminów', 1)
    for item in [
        'W dokumentach projektu używać nazwy SWDsuw, rozwijanej jako „System wspomagania dyplomowania studentów uczelni wyższej”. Skrót „SWD” z nazwy zadania oznacza ten sam projekt, jeśli Zamawiający nie wskaże inaczej.',
        'Rozróżniać „pracę dyplomową” (obiekt) od „tematu pracy” (przedmiot wyboru i przypisania promotora).',
        'Rozróżniać ocenę pracy jako całości od oceny przypisanej indywidualnemu autorowi.',
        'Nie używać „ocena końcowa” zamiennie z „wynikiem ukończenia studiów”; algorytm wyniku jest zewnętrzną regułą statutową.',
        'Tekst Soft-AS/SWS stanowi oddzielne ćwiczenie analizy lingwistycznej i nie jest źródłem terminów SWDsuw.',
    ]:
        bullet(doc, item)
    return doc


TASKS = [
    ('1.1', 'Inception', 'Project Management', 'Kick-off, potwierdzenie zakresu i kanału decyzji', 'Scrum Master / Project Manager; przedstawiciel Zamawiającego', 1, 1, '', 'Uzgodniony cel, interesariusze, sposób raportowania'),
    ('1.2', 'Inception', 'Requirements', 'Analiza źródła PZ SWD i szablonów; rejestr faktów, założeń i pytań', 'Analityk systemowy; Zamawiający', 1, 2, '1.1', 'Rejestr źródeł i kwestii do decyzji'),
    ('1.3', 'Inception', 'Business Modeling', 'Rozpoznanie aktorów biznesowych i procesów dyplomowania', 'Analityk procesów; promotor / dziekanat', 1, 2, '1.2', 'Mapa aktorów i procesów'),
    ('1.4', 'Inception', 'Requirements', 'Przedstawienie problemu (Zamawiający i Wykonawca) oraz wizja produktu', 'Analityk wymagań; Scrum Master; Zamawiający', 1, 2, '1.2, 1.3', 'Zatwierdzony problem i wizja — M1'),
    ('1.5', 'Inception', 'Business Modeling', 'Analiza lingwistyczna niezależnego tekstu Soft-AS/SWS', 'Analityk domenowy', 2, 1, '1.1', 'Pokolorowany tekst, tabele grup, wynik kandydatów'),
    ('1.6', 'Inception', 'Requirements', 'Słownik domenowy SWDsuw i przegląd z Zamawiającym', 'Analityk domenowy; promotor / dziekanat', 2, 1, '1.2, 1.3', 'Uzgodnione nazwy pojęć'),
    ('1.7', 'Inception', 'Project Management', 'Ocena ryzyk, plan faz/iteracji, zasoby i raport Scrum Mastera', 'Scrum Master / Project Manager; Project Reviewer', 2, 1, '1.4, 1.5, 1.6', 'Wstępny baseline; przegląd fazy — M2'),
    ('2.1', 'Elaboration', 'Requirements', 'Rozwinięcie wymagań i szkiców przypadków użycia; śledzenie źródeł', 'Analityk wymagań; specyfikator przypadków użycia', 3, 2, '1.7', 'Bazowy zakres przypadków użycia'),
    ('2.2', 'Elaboration', 'Analysis & Design', 'Architektura rozwiązania, model danych domeny, wariant integracji „Plagiat”', 'Architekt; analityk; przedstawiciel systemu zewnętrznego', 3, 3, '2.1', 'Prototyp/architektura do oceny — M3'),
    ('2.3', 'Elaboration', 'Test', 'Plan testów akceptacyjnych, bezpieczeństwa, śledzenia i odbioru', 'Test Manager; tester; Zamawiający', 4, 2, '2.1, 2.2', 'Kryteria odbioru i podejście testowe'),
    ('2.4', 'Elaboration', 'Configuration & Change Management', 'Zasady wersjonowania, zarządzania zmianą i baseline wymagań', 'Configuration Manager; Scrum Master', 4, 1, '2.1', 'Plan zarządzania zmianą'),
    ('2.5', 'Elaboration', 'Project Management', 'Przegląd architektury, ryzyk, zakresu i planu iteracji', 'Project Reviewer; Scrum Master; Zamawiający', 5, 1, '2.2, 2.3, 2.4', 'Zatwierdzona baza do budowy — M4'),
    ('3.1', 'Construction', 'Iteration C1', 'Iteracja 1: rejestracja prac, autorów, tematów, promotorów i powiązań', 'Designer; Developer Role Set; Tester', 6, 2, '2.5', 'Przegląd iteracji C1'),
    ('3.2', 'Construction', 'Iteration C1', 'Testy integracyjne i weryfikacja indywidualnych autorów/tematów', 'Test Analyst; Tester; Requirements Reviewer', 7, 1, '3.1', 'Akceptacja iteracji C1 — M5'),
    ('3.3', 'Construction', 'Iteration C2', 'Iteracja 2: dokumenty pracy, przekazanie do antyplagiatu, recenzje i oceny', 'Developer Role Set; integrator; Tester', 8, 2, '3.2', 'Przegląd iteracji C2'),
    ('3.4', 'Construction', 'Iteration C2', 'Testy uprawnień, błędów integracji i ocen per autor', 'Test Manager; Test Designer; Security Reviewer', 9, 1, '3.3', 'Akceptacja iteracji C2 — M6'),
    ('3.5', 'Construction', 'Iteration C3', 'Iteracja 3: komisje, pytania, protokół obrony, wynik i raporty', 'Developer Role Set; Business Designer; Tester', 10, 2, '3.4', 'Przegląd iteracji C3'),
    ('3.6', 'Construction', 'Test', 'Testy regresji, wyszukiwania, raportów i śledzenia wymagań', 'Tester; Test Analyst; Project Reviewer', 11, 1, '3.5', 'Zakres funkcji gotowy do odbioru — M7'),
    ('3.7', 'Construction', 'Configuration & Change Management', 'Ciągłe buildy, audyt zmian, baseline i kontrola defektów', 'Configuration Manager; Developer Role Set; Tester', 6, 6, '2.5', 'Status zmian i baseline per iteracja'),
    ('4.1', 'Transition', 'Test', 'Testy akceptacyjne z użytkownikami i zamknięcie defektów blokujących', 'Zamawiający; Test Manager; Scrum Master', 12, 1, '3.6', 'UAT i decyzja o dopuszczeniu — M8'),
    ('4.2', 'Transition', 'Deployment', 'Przygotowanie wdrożenia, instrukcji, szkoleń i ewentualnej migracji', 'Deployment Manager; Technical Writer; administrator', 12, 2, '4.1', 'Pakiet wdrożeniowy i szkolenie'),
    ('4.3', 'Transition', 'Project Management', 'Odbiór, przegląd końcowy, raport zamknięcia przesyłany przez Scrum Mastera', 'Scrum Master; Project Reviewer; Zamawiający', 13, 1, '4.2', 'Odbiór produktu i zamknięcie — M9'),
]

RESOURCES = [
    ('Scrum Master / Project Manager', 'Planowanie, raportowanie statusu, ryzyka, zależności i koordynacja; wyłącznie Scrum Master przekazuje raport.', 'Inception–Transition', 'Nazwisko i dostępność do uzupełnienia'),
    ('Przedstawiciel Zamawiającego / sponsor', 'Decyzje dotyczące polityk uczelni, priorytetów, zakresu i odbioru.', 'Inception, Elaboration, UAT', 'Nazwa uczelni i osoby do potwierdzenia'),
    ('Analityk systemowy / wymagań', 'Problemy, potrzeby, słownik, śledzenie wymagań, procesy i przypadki użycia.', 'Inception–Elaboration', 'Rola z puli RUP'),
    ('Analityk procesów biznesowych', 'Modelowanie procesu dyplomowania i terminologii.', 'Inception–Elaboration', 'Rola z puli RUP'),
    ('Architekt oprogramowania', 'Granica systemu, architektura i wykonalność integracji z antyplagiatem.', 'Elaboration', 'Rola z puli RUP'),
    ('Specyfikator przypadków użycia / Designer', 'Rozwinięcie celów aktorów i projektów interakcji domenowych.', 'Elaboration–Construction', 'Rola z puli RUP'),
    ('Developer Role Set', 'Implementacja komponentów, integracja i poprawki.', 'Construction–Transition', 'Zespół i dostępność niepodane'),
    ('Test Manager / Test Analyst / Tester', 'Strategia testów, scenariusze, testy integracji, regresji i akceptacji.', 'Elaboration–Transition', 'Rola z puli RUP'),
    ('Configuration Manager / Change Control Manager', 'Baseline, wersjonowanie, audyt i obsługa zmian.', 'Wszystkie fazy', 'Rola z puli RUP'),
    ('Technical Writer / Deployment Manager', 'Instrukcje, materiały szkoleniowe i pakiet przejścia.', 'Transition', 'Rola z puli RUP'),
    ('Administrator / przedstawiciel systemu „Plagiat”', 'Warunki dostępu, integracja, polityka retencji i obsługa błędów.', 'Elaboration–Construction', 'Osoba/organizacja do wskazania'),
]

MILESTONES = [
    ('M1', 'Problem, wizja i interesariusze zatwierdzeni', 'Koniec T2', 'Zamawiający + Scrum Master'),
    ('M2', 'Baseline fazy Inception i wstępny plan iteracji', 'Koniec T2', 'Przegląd projektu; raport wysyła Scrum Master'),
    ('M3', 'Wykonalność architektury i integracji oceniona', 'Koniec T5', 'Project Reviewer + właściciel systemu antyplagiatowego'),
    ('M4', 'Zakres wymagań i podejście testowe zaakceptowane', 'Koniec T5', 'Zamawiający'),
    ('M5', 'Iteracja C1 oceniona', 'Koniec T7', 'Test Manager + Zamawiający'),
    ('M6', 'Iteracja C2 oceniona', 'Koniec T9', 'Test Manager + Zamawiający'),
    ('M7', 'Funkcje Construction gotowe do odbioru', 'Koniec T11', 'Project Reviewer'),
    ('M8', 'Testy akceptacyjne zakończone', 'Koniec T12', 'Użytkownicy reprezentatywni + Zamawiający'),
    ('M9', 'Odbiór / przekazanie i raport zamknięcia', 'Koniec T13', 'Zamawiający; raport przesyła wyłącznie Scrum Master'),
]


def build_schedule_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Gantt'
    headers = ['WBS', 'Faza RUP', 'Dyscyplina / iteracja', 'Zadanie', 'Odpowiedzialna rola', 'Start (tydz.)', 'Czas (tyg.)', 'Poprzedniki', 'Produkt / kryterium'] + [f'T{i}' for i in range(1, 14)]
    ws.append(headers)
    for row in TASKS:
        wbs, phase, discipline, name, role, start, duration, predecessor, output = row
        values = list(row) + [''] * 13
        for week in range(start, start + duration):
            if 1 <= week <= 13:
                values[8 + week] = '■'
        ws.append(values)
    header_fill = PatternFill('solid', fgColor=NAVY)
    thin = Side(style='thin', color='D8DEE6')
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF', size=9)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    phase_colors = {'Inception': BLUE, 'Elaboration': PURPLE, 'Construction': GREEN, 'Transition': ORANGE}
    for row_idx in range(2, ws.max_row + 1):
        phase = ws.cell(row_idx, 2).value
        fill = PatternFill('solid', fgColor=phase_colors[phase])
        for col in range(1, 10):
            c = ws.cell(row_idx, col)
            c.fill = fill if col <= 2 else PatternFill('solid', fgColor='FFFFFF' if row_idx % 2 == 0 else 'F5F7F9')
            c.alignment = Alignment(vertical='center', wrap_text=True)
            c.border = Border(bottom=thin)
            c.font = Font(size=8.5, bold=(col in (1, 2)))
        for col in range(10, 23):
            c = ws.cell(row_idx, col)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = Border(bottom=thin, right=thin)
            if c.value == '■':
                c.fill = fill
                c.font = Font(color=NAVY, bold=True, size=11)
    widths = {1: 8, 2: 15, 3: 24, 4: 48, 5: 36, 6: 12, 7: 11, 8: 18, 9: 48}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for col in range(10, 23):
        ws.column_dimensions[get_column_letter(col)].width = 5
    ws.freeze_panes = 'J2'
    ws.auto_filter.ref = f'A1:V{ws.max_row}'
    ws.row_dimensions[1].height = 34
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 34
    ws.sheet_view.showGridLines = False

    res = wb.create_sheet('Zasoby')
    res.append(['Rola / zasób RUP', 'Odpowiedzialność', 'Faza / okres', 'Uwagi / obsada'])
    for row in RESOURCES:
        res.append(row)
    for cell in res[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for row in res.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(bottom=thin)
    for col, width in zip('ABCD', [38, 75, 28, 42]):
        res.column_dimensions[col].width = width
    res.freeze_panes = 'A2'
    res.sheet_view.showGridLines = False

    ms = wb.create_sheet('Kamienie milowe')
    ms.append(['ID', 'Kamień milowy / decyzja', 'Termin względny', 'Odpowiedzialność / dowód'])
    for row in MILESTONES:
        ms.append(row)
    for cell in ms[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
    for row in ms.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(bottom=thin)
    for col, width in zip('ABCD', [10, 56, 18, 56]):
        ms.column_dimensions[col].width = width
    ms.freeze_panes = 'A2'
    ms.sheet_view.showGridLines = False

    assumptions = wb.create_sheet('Założenia')
    assumptions.append(['Założenie / ograniczenie', 'Opis'])
    for row in [
        ('Horyzont', 'Wstępny baseline 13 tygodni od zatwierdzonego kick-offu; kolumny T1–T13 są tygodniami względnymi, nie datami kalendarzowymi.'),
        ('RUP', 'Inception T1–T2; Elaboration T3–T5; Construction T6–T11 (3 iteracje); Transition T12–T13.'),
        ('Obsada', 'Użyto nazw ról z szablonu RUP_Project.mpp; nazwiska, wielkość zespołu i dostępność nie są podane.'),
        ('Szablon MPP', 'Plik RUP_Project.mpp zawiera przykładowy, datowany na 2005 r. plan (690 zadań); jego dat i szczegółowej siatki zadań nie przeniesiono jako baseline SWDsuw.'),
        ('Raport', 'Status/raport do SharePoint wysyła wyłącznie Scrum Master; brak adresu witryny i uprawnień, więc pliku nie opublikowano.'),
        ('Zakres', 'Plan ma charakter wstępny; należy go zaktualizować po potwierdzeniu składu zespołu, dostępności właścicieli procesów, terminów akademickich i integracji „Plagiat”.'),
        ('Checklist Excel', 'Lista_kontrolna.xls zawiera kontrolę jakości przypadków użycia (LK wym1/wym2), a nie checklistę wszystkich sześciu artefaktów. Kryteria wykorzystano jako bramkę jakości dla przyszłych opisów przypadków użycia.'),
    ]:
        assumptions.append(row)
    for cell in assumptions[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
    for row in assumptions.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    assumptions.column_dimensions['A'].width = 28
    assumptions.column_dimensions['B'].width = 120
    assumptions.sheet_view.showGridLines = False

    out = ROOT / '04_Wstepny_harmonogram_RUP_SWDSUW.xlsx'
    wb.save(out)
    return out


def build_gantt_pdf():
    out = ROOT / '04_Wstepny_harmonogram_RUP_SWDSUW.pdf'
    page = landscape(A3)
    doc = SimpleDocTemplate(str(out), pagesize=page, rightMargin=12 * mm, leftMargin=12 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle('TitlePL', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=colors.HexColor('#1F4E78'), spaceAfter=5)
    body = ParagraphStyle('BodyPL', parent=styles['BodyText'], fontName='Helvetica', fontSize=7.5, leading=9)
    small = ParagraphStyle('SmallPL', parent=body, fontSize=6.5, leading=7.2)
    story = [Paragraph('Wstępny harmonogram RUP — SWDsuw', title), Paragraph('Baseline 13 tygodni względnych od kick-offu; do zatwierdzenia przez Zamawiającego. Szczegóły, zasoby i założenia znajdują się w skoroszycie XLSX.', body), Spacer(1, 5 * mm)]
    headers = ['WBS', 'Faza', 'Zadanie / produkt', 'Rola wiodąca', 'Start', 'Czas'] + [f'T{i}' for i in range(1, 14)]
    data = [[Paragraph(x, small) for x in headers]]
    phase_color = {'Inception': colors.HexColor('#D9EAF7'), 'Elaboration': colors.HexColor('#E4DFEC'), 'Construction': colors.HexColor('#E2F0D9'), 'Transition': colors.HexColor('#FCE4D6')}
    for wbs, phase, discipline, name, role, start, duration, prev, output in TASKS:
        row = [wbs, phase, name, role, f'T{start}', f'{duration} t.']
        row.extend(['■' if start <= week < start + duration else '' for week in range(1, 14)])
        data.append([Paragraph(str(value), small) for value in row])
    widths = [11*mm, 22*mm, 58*mm, 50*mm, 13*mm, 13*mm] + [7*mm]*13
    gantt = Table(data, colWidths=widths, repeatRows=1)
    command = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (6, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#C7D1DB')),
        ('LEFTPADDING', (0, 0), (-1, -1), 2), ('RIGHTPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]
    for idx, task in enumerate(TASKS, start=1):
        fill = phase_color[task[1]]
        command.append(('BACKGROUND', (0, idx), (1, idx), fill))
        for week in range(task[5], task[5] + task[6]):
            if 1 <= week <= 13:
                command.append(('BACKGROUND', (5 + week, idx), (5 + week, idx), fill))
                command.append(('TEXTCOLOR', (5 + week, idx), (5 + week, idx), colors.HexColor('#1F4E78')))
                command.append(('FONTNAME', (5 + week, idx), (5 + week, idx), 'Helvetica-Bold'))
    gantt.setStyle(TableStyle(command))
    story.append(gantt)
    story.extend([Spacer(1, 4 * mm), Paragraph('Kamienie milowe', styles['Heading2'])])
    ms_data = [['ID', 'Kamień milowy', 'Termin', 'Odpowiedzialność / dowód']]
    ms_data.extend([list(row) for row in MILESTONES])
    milestone_table = Table([[Paragraph(str(x), small) for x in row] for row in ms_data], colWidths=[12*mm, 78*mm, 22*mm, 90*mm], repeatRows=1)
    milestone_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#C7D1DB')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 3), ('RIGHTPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(milestone_table)
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph('Raport ze stanu projektu i jego przekazanie do SharePoint należy wyłącznie do Scrum Mastera. Daty kalendarzowe, nazwiska i dostępność są do uzupełnienia.', small))
    doc.build(story)
    return out


def build_readme():
    text = '''# Dokumentacja Etapu A — SWDsuw

Data opracowania: 21.10.2026. Wersje DOCX są edytowalne; PDF to kopie do przeglądu i przesyłania.

## Artefakty

1. `01_Przedstawienie_problemu_Zamawiajacego.docx/.pdf` — uporządkowany zapis treści z `PZ SWD przedstawienie problemu.docx`.
2. `02_Przedstawienie_problemu_Wykonawcy_SWDSUW.docx/.pdf` — analiza problemu, aktorów, procesów, granicy i kwestii do uzgodnienia.
3. `03_Wizja_projektu_SWDSUW.docx/.pdf` — wypełnienie sekcji szablonu wizji projektu.
4. `04_Wstepny_harmonogram_RUP_SWDSUW.xlsx/.pdf` — wykres Gantta, zadania, role, zależności, kamienie milowe i założenia.
5. `05_Analiza_lingwistyczna_SWS_Soft-AS.docx/.pdf` — kolorowany tekst Soft-AS/SWS, osobne tabele grup czasownikowych i rzeczownikowych oraz wyprowadzone opisowo kandydatury obiektów.
6. `06_Slownik_pojec_SWDSUW.docx/.pdf` — słownik pojęć domenowych SWDsuw; bez technologii.

## Granica zakresu

Źródłowy dokument projektu opisuje SWDsuw (dyplomowanie). Tekst Soft-AS/SWS został dostarczony jako materiał do dodatkowej analizy lingwistycznej, dlatego jest osobnym artefaktem i jego obiekty nie zostały przeniesione do słownika SWDsuw. MSIoPS wykorzystano wyłącznie jako przykład formatu.

## Założenia i otwarte decyzje

- Nazwę SWDsuw przyjęto z dokumentu źródłowego; skrót „SWD” w nazwie zadania traktowany jest jako odniesienie do tego systemu.
- Gantt ma horyzont 13 tygodni względnych od zatwierdzonego kick-offu. Nie podano daty startu, składu ani dostępności osób, więc użyto nazw ról.
- `RUP_Project.mpp` zawiera przykładowy harmonogram z 2005 r. i role RUP; jego historycznych dat oraz szablonowych 690 pozycji nie skopiowano jako planu SWDsuw.
- `Lista_kontrolna.xls` to checklista jakości przypadków użycia (31 kryteriów), a nie checklista kompletności tych sześciu dokumentów.
- Zgodnie z instrukcją tylko Scrum Master przekazuje raport do SharePoint. Nie podano witryny ani uprawnień, więc materiały są gotowe lokalnie i nie zostały opublikowane.
- Nie wpisano nazwisk autorów, uczelni ani niepotwierdzonych rozwiązań technicznych.

## Rejestr kwestii do zatwierdzenia

Kluczowe pytania o kryterium antyplagiatowe, algorytm oceny końcowej, liczbę i ocenianie pytań na obronie, uprawnienia, recenzentów, raporty oraz retencję są w dokumentach A2 i A3. Są to decyzje Zamawiającego, a nie założenia Wykonawcy.
'''
    (ROOT / '00_README_artefakty.md').write_text(text, encoding='utf-8')


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    build_client_problem().save(ROOT / '01_Przedstawienie_problemu_Zamawiajacego.docx')
    build_contractor_problem().save(ROOT / '02_Przedstawienie_problemu_Wykonawcy_SWDSUW.docx')
    build_vision().save(ROOT / '03_Wizja_projektu_SWDSUW.docx')
    build_linguistic_analysis().save(ROOT / '05_Analiza_lingwistyczna_SWS_Soft-AS.docx')
    build_glossary().save(ROOT / '06_Slownik_pojec_SWDSUW.docx')
    build_schedule_xlsx()
    build_gantt_pdf()
    build_readme()
    print(f'Wygenerowano artefakty w: {ROOT}')


if __name__ == '__main__':
    main()
